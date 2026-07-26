#!/usr/bin/env python3
"""Genera il PPTX SAL dal Markdown sorgente, rispettando il modello aziendale.

La source of truth e' ``presentazioni.md``. Il file PPTX indicato in
``DEFAULT_TEMPLATE`` e' usato esclusivamente come modello grafico: la slide 1
e le ultime tre slide sono conservate; tutte le slide intermedie sono eliminate
e sostituite con nuove slide create dal layout ``Titolo e contenuto``.

Requisiti:
    py -3 -m pip install python-pptx

Uso:
    py -3 genera_pptx_sal.py
    py -3 genera_pptx_sal.py --output "C:\\percorso\\SAL-generato.pptx"
    py -3 genera_pptx_sal.py --check
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Pt


DEFAULT_TEMPLATE = Path(
    r"C:\Users\alexandru\OneDrive - IRION\Analisi\test estensioni\SAL-27072026.pptx"
)
DEFAULT_OUTPUT = DEFAULT_TEMPLATE.with_name(f"{DEFAULT_TEMPLATE.stem}-draft.pptx")
SOURCE_MARKDOWN = Path(__file__).with_name("presentazioni.md")
CONTENT_LAYOUT_NAME = "Titolo e contenuto "
PRESERVED_FINAL_SLIDES = 3
TOPIC_LABEL = "SAL · DuckDB ed estensioni"


@dataclass(frozen=True)
class SlideSpec:
    number: int
    title: str
    body: tuple[str, ...]


def clean_inline_markdown(value: str) -> str:
    """Conserva il testo, eliminando soltanto la sintassi Markdown non visibile."""
    value = re.sub(r"!\[[^]]*\]\([^)]*\)", "", value)
    value = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", value)
    value = value.replace("**", "").replace("__", "").replace("`", "")
    return re.sub(r"\s+", " ", value).strip()


def markdown_to_paragraphs(markdown: str) -> tuple[str, ...]:
    """Trasforma il corpo di una slide in paragrafi leggibili e non in HTML."""
    paragraphs: list[str] = []
    current: list[str] = []

    def flush_current() -> None:
        if current:
            text = clean_inline_markdown(" ".join(current))
            if text:
                paragraphs.append(text)
            current.clear()

    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            flush_current()
            continue
        if line == "---" or re.match(r"^!\[[^]]*\]\([^)]*\)$", line):
            flush_current()
            continue
        if line.startswith("### ") or line.startswith("#### "):
            flush_current()
            heading = clean_inline_markdown(re.sub(r"^#{3,4}\s+", "", line))
            if heading:
                paragraphs.append(heading)
            continue

        bullet = re.match(r"^(?:[-*+]\s+|\d+[.)]\s+)(.+)$", line)
        if bullet:
            flush_current()
            paragraphs.append("• " + clean_inline_markdown(bullet.group(1)))
            continue
        if line.startswith("> "):
            flush_current()
            paragraphs.append(clean_inline_markdown(line[2:]))
            continue
        current.append(line)

    flush_current()
    return tuple(paragraphs)


def parse_source(markdown_path: Path) -> list[SlideSpec]:
    """Legge esclusivamente le sezioni `## Slide N` della source of truth."""
    if not markdown_path.is_file():
        raise FileNotFoundError(f"Source Markdown non trovata: {markdown_path}")

    text = markdown_path.read_text(encoding="utf-8")
    starts = list(re.finditer(r"(?m)^##\s+Slide\s+(\d+)\b[^\n]*$", text))
    if not starts:
        raise ValueError("Nessuna sezione `## Slide N` trovata nella source Markdown.")

    specs: list[SlideSpec] = []
    for index, match in enumerate(starts):
        block_end = starts[index + 1].start() if index + 1 < len(starts) else len(text)
        block = text[match.end() : block_end]
        title_match = re.search(r"(?m)^#\s+(.+?)\s*$", block)
        if not title_match:
            raise ValueError(f"La slide {match.group(1)} non contiene un titolo `#`.")
        title = clean_inline_markdown(title_match.group(1))
        body = markdown_to_paragraphs(block[title_match.end() :])
        if not title:
            raise ValueError(f"La slide {match.group(1)} ha un titolo vuoto.")
        if len(title) > 120:
            raise ValueError(
                f"Il titolo della slide {match.group(1)} e' troppo lungo ({len(title)} caratteri). "
                "Accorcialo nella source of truth anziche' ridurre il font."
            )
        if len(body) > 14 or sum(len(item) for item in body) > 1_550:
            raise ValueError(
                f"Il contenuto della slide {match.group(1)} e' troppo denso. "
                "Dividilo in piu' sezioni `## Slide N` nel Markdown."
            )
        specs.append(SlideSpec(int(match.group(1)), title, body))

    expected_numbers = list(range(1, len(specs) + 1))
    actual_numbers = [spec.number for spec in specs]
    if actual_numbers != expected_numbers:
        raise ValueError(
            f"Le slide devono essere numerate consecutivamente da 1: trovate {actual_numbers}."
        )
    return specs


def find_content_layout(presentation: Presentation):
    for layout in presentation.slide_layouts:
        if layout.name == CONTENT_LAYOUT_NAME:
            return layout
    raise ValueError(
        f"Il modello non contiene il layout {CONTENT_LAYOUT_NAME!r}; non posso rispettarne la grafica."
    )


def remove_slide(presentation: Presentation, index: int) -> None:
    """Rimuove una slide dalla sequenza senza alterare le slide preservate.

    Non si elimina la relazione OOXML: ``python-pptx`` puo' riutilizzare il
    medesimo nome di parte alla successiva ``add_slide`` e produrre un archivio
    PPTX con XML duplicato. La relazione orfana non e' raggiungibile dalla
    sequenza di slide e quindi non e' visibile nel deck generato.
    """
    slide_id = presentation.slides._sldIdLst[index]
    del presentation.slides._sldIdLst[index]


def set_text(shape, value: str, *, font_size: int | None = None) -> None:
    text_frame = shape.text_frame
    text_frame.clear()
    paragraph = text_frame.paragraphs[0]
    paragraph.text = value
    if font_size:
        for run in paragraph.runs:
            run.font.size = Pt(font_size)


def placeholders_by_type(slide, placeholder_type: PP_PLACEHOLDER) -> list:
    return [
        shape
        for shape in slide.placeholders
        if shape.placeholder_format.type == placeholder_type
    ]


def populate_content_slide(slide, spec: SlideSpec) -> None:
    titles = placeholders_by_type(slide, PP_PLACEHOLDER.TITLE)
    bodies = placeholders_by_type(slide, PP_PLACEHOLDER.BODY)
    if not titles or len(bodies) < 2:
        raise ValueError(
            "Il layout del modello non espone i placeholder titolo, etichetta e contenuto attesi."
        )

    set_text(titles[0], spec.title)
    set_text(bodies[0], TOPIC_LABEL)

    content = bodies[-1]
    content.text_frame.clear()
    for index, value in enumerate(spec.body):
        paragraph = content.text_frame.paragraphs[0] if index == 0 else content.text_frame.add_paragraph()
        paragraph.text = value
        paragraph.level = 0
        for run in paragraph.runs:
            run.font.size = Pt(17)


def move_new_slides_after_title(presentation: Presentation, new_slide_count: int) -> None:
    """`python-pptx` aggiunge in coda; questa funzione ricolloca le nuove slide alla posizione 2."""
    slide_ids = presentation.slides._sldIdLst
    new_ids = list(slide_ids)[-new_slide_count:]
    for slide_id in new_ids:
        slide_ids.remove(slide_id)
    for offset, slide_id in enumerate(new_ids):
        slide_ids.insert(1 + offset, slide_id)


def generate(template: Path, markdown: Path, output: Path) -> None:
    if not template.is_file():
        raise FileNotFoundError(f"Modello PPTX non trovato: {template}")

    specs = parse_source(markdown)
    if len(specs) < 2:
        raise ValueError("La source deve contenere almeno Slide 1 e Slide 2.")

    presentation = Presentation(template)
    if len(presentation.slides) < PRESERVED_FINAL_SLIDES + 1:
        raise ValueError("Il modello deve contenere almeno la slide iniziale e le ultime tre da preservare.")

    original_slide_count = len(presentation.slides)
    # Le nuove slide vanno create prima di rimuovere quelle sostituite. In questo
    # modo python-pptx assegna nuovi nomi OOXML (slide6.xml, slide7.xml, ...)
    # senza collidere con le tre slide finali da preservare.
    content_layout = find_content_layout(presentation)
    for spec in specs[1:]:  # Slide 1 e' la cover del modello: resta invariata.
        new_slide = presentation.slides.add_slide(content_layout)
        populate_content_slide(new_slide, spec)

    # Elimina soltanto il segmento che va dalla slide 2 alla terzultima slide esclusa.
    last_replaced_index = original_slide_count - PRESERVED_FINAL_SLIDES - 1
    for index in range(last_replaced_index, 0, -1):
        remove_slide(presentation, index)
    move_new_slides_after_title(presentation, len(specs) - 1)

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    print(f"PPTX generato: {output}")
    print(
        f"Slide totali: {len(presentation.slides)} "
        f"(1 cover preservata, {len(specs) - 1} generate dal Markdown, "
        f"{PRESERVED_FINAL_SLIDES} finali preservate)."
    )


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--source", type=Path, default=SOURCE_MARKDOWN)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--check",
        action="store_true",
        help="Convalida la source Markdown senza creare o modificare alcun PPTX.",
    )
    args = parser.parse_args(argv)

    try:
        specs = parse_source(args.source)
        if args.check:
            print(f"Source valida: {args.source} ({len(specs)} slide dichiarate).")
            return 0
        generate(args.template, args.source, args.output)
        return 0
    except (FileNotFoundError, ValueError) as error:
        print(f"Errore: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
